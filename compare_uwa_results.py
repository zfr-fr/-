import argparse
import csv
import json
import os
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional, Tuple

try:
    import openpyxl
except ImportError:
    print("请先安装依赖: pip install openpyxl")
    raise


@dataclass
class ViolationRecord:
    file: str
    line: Optional[int]
    rule: str
    description: str
    source: str
    sheet: str = ""
    row: Optional[int] = None
    line_content: str = ""
    canon_rule: str = ""


def normalize_path(path: str) -> str:
    if not path:
        return ""
    text = str(path).strip().replace("\\", "/")
    text = re.sub(r"^[a-zA-Z]:/", "", text)
    text = re.sub(r"/+", "/", text)
    return text.lower()


def normalize_rule(rule: str) -> str:
    if not rule:
        return ""
    return re.sub(r"\s+", "", str(rule).strip().upper())


def normalize_desc(desc: str) -> str:
    if not desc:
        return ""
    text = str(desc).strip().lower()
    text = re.sub(r"\s+", "", text)
    text = re.sub(r"[，。,;；:：\-_/()（）\[\]{}<>\"'`~!@#$%^&*+=|\\]", "", text)
    return text


def parse_line_number(value) -> Optional[int]:
    if value is None:
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    match = re.search(r"\d+", str(value))
    if match:
        return int(match.group(0))
    return None


def extract_line_numbers(text: str) -> List[int]:
    if not text:
        return []
    tokens = re.split(r"[,\s，]+", text)
    lines: List[int] = []
    for token in tokens:
        token = token.strip()
        if not token:
            continue
        if re.match(r"^\d+$", token):
            lines.append(int(token))
            continue
        range_match = re.match(r"^(\d+)\s*[-~]\s*(\d+)$", token)
        if range_match:
            start = int(range_match.group(1))
            end = int(range_match.group(2))
            if start <= end:
                lines.extend(range(start, end + 1))
            else:
                lines.extend(range(end, start + 1))
            continue
        maybe = parse_line_number(token)
        if maybe is not None:
            lines.append(maybe)
    return lines


def parse_source_location(
    value: str,
    rule: str,
    desc: str,
    source_file: str,
    sheet: str,
    row: int,
) -> List[ViolationRecord]:
    if not value:
        return []
    text = str(value).strip()
    if not text:
        return []

    records: List[ViolationRecord] = []
    parts = re.split(r"[|\n]+", text)
    for part in parts:
        segment = part.strip()
        if not segment:
            continue
        segment = segment.split(";")[0].strip()

        match = re.search(r"(?P<path>.*?\.cs)\s*:\s*(?P<lines>.+)", segment, re.IGNORECASE)
        if not match:
            continue

        file_path = match.group("path").strip()
        lines_text = match.group("lines").strip()
        lines = extract_line_numbers(lines_text)
        if not lines:
            records.append(
                ViolationRecord(
                    file=file_path,
                    line=None,
                    rule=rule,
                    description=desc,
                    source=source_file,
                    sheet=sheet,
                    row=row,
                )
            )
            continue

        for line in lines:
            records.append(
                ViolationRecord(
                    file=file_path,
                    line=line,
                    rule=rule,
                    description=desc,
                    source=source_file,
                    sheet=sheet,
                    row=row,
                )
            )
    return records


def matches_target_path(file_path: str, target_norm: str) -> bool:
    if not target_norm:
        return True
    file_norm = normalize_path(file_path)
    if target_norm in file_norm:
        return True
    assets_index = target_norm.find("assets/")
    if assets_index >= 0:
        assets_suffix = target_norm[assets_index:]
        if assets_suffix in file_norm:
            return True
    return False


def filter_records_by_path(records: List[ViolationRecord], target_path: Optional[str]) -> List[ViolationRecord]:
    if not target_path:
        return records
    target_norm = normalize_path(target_path)
    return [record for record in records if matches_target_path(record.file, target_norm)]


def load_rule_map(path: Optional[str]) -> Dict[str, str]:
    if not path:
        return {}
    with open(path, "r", encoding="utf-8") as f:
        data = json.load(f)
    return {normalize_rule(k): normalize_rule(v) for k, v in data.items()}


def load_rule_excel_map(path: Optional[str]) -> Dict[str, Dict[str, List[str]]]:
    if not path:
        return {}
    with open(path, "r", encoding="utf-8") as f:
        data = json.load(f)
    normalized: Dict[str, Dict[str, List[str]]] = {}
    for rule_key, value in data.items():
        canon_rule = normalize_rule(rule_key)
        entry: Dict[str, List[str]] = {"keywords": [], "regex": []}
        if isinstance(value, str):
            entry["keywords"] = [value.lower()]
        elif isinstance(value, list):
            entry["keywords"] = [str(item).lower() for item in value if item]
        elif isinstance(value, dict):
            keywords = value.get("keywords", []) or value.get("keyword", [])
            regex = value.get("regex", []) or value.get("pattern", [])
            if isinstance(keywords, str):
                keywords = [keywords]
            if isinstance(regex, str):
                regex = [regex]
            entry["keywords"] = [str(item).lower() for item in keywords if item]
            entry["regex"] = [str(item) for item in regex if item]
        normalized[canon_rule] = entry
    return normalized


def default_rule_excel_map() -> Dict[str, Dict[str, List[str]]]:
    return {
        "DEBUG_LOG_USAGE": {
            "keywords": ["unityengine.debug类的log函数", "debug类的log函数"],
        },
        "NEW_IN_UPDATE_METHOD": {
            "keywords": ["update中存在引用类型的new操作"],
        },
        "EMPTY_UPDATE_METHOD": {
            "keywords": ["空的update", "lateupdate", "fixedupdate方法的检测数据"],
        },
        "FIND_FUNCTION_CALLS": {
            "keywords": ["find类函数调用"],
        },
        "FOREACH_ON_SPECIFIC_CONTAINERS": {
            "keywords": ["特定容器的foreach遍历"],
        },
        "GET_COMPONENTS_IN_CHILDREN": {
            "keywords": ["getcomponentsinchildren调用"],
        },
        "GET_COMPONENTS_IN_PARENT": {
            "keywords": ["getcomponentsinparent调用"],
        },
        "COMPUTE_BUFFER_GETDATA": {
            "keywords": ["computebuffer.getdata调用"],
        },
        "TEXTURE_GETPIXELS_CALL": {
            "keywords": ["getpixels32", "getpixels", "setpixels_getpixels32"],
        },
        "TEXTURE_SETPIXELS_CALL": {
            "keywords": ["setpixels", "setpixels_getpixels32"],
        },
        "TEXTASSET_BYTES_USAGE": {
            "keywords": ["textasset.www.bytes调用", "textasset.bytes调用"],
        },
        "CAMERA_MAIN_USAGE": {
            "keywords": ["camera.main的调用"],
        },
        "RENDERER_MATERIAL_GETTER": {
            "keywords": ["material_materials的获取"],
        },
        "RENDERER_SHAREDMATERIALS_GETTER": {
            "keywords": ["sharedmaterials的获取"],
        },
        "ONGUI_METHOD_PRESENT": {
            "keywords": ["ongui方法的检测数据"],
        },
        "SENDMESSAGE_USAGE": {
            "keywords": ["gameobject.sendmessage调用"],
        },
        "TAG_PROPERTY_USAGE": {
            "keywords": [".tag的调用"],
        },
        "LINQ_USAGE": {
            "keywords": ["linq相关函数的调用"],
        },
        "REFLECTION_USAGE": {
            "keywords": ["reflection相关函数的调用"],
        },
        "HEAP_ALLOCATING_STRING_OPS": {
            "keywords": ["堆内存分配的字符串操作"],
        },
    }


def rule_from_excel_filename(
    filename: str, rule_excel_map: Dict[str, Dict[str, List[str]]]
) -> List[str]:
    if not filename or not rule_excel_map:
        return []
    name = os.path.basename(filename).lower()
    matched: List[str] = []
    for rule, spec in rule_excel_map.items():
        for keyword in spec.get("keywords", []):
            if keyword and keyword in name:
                matched.append(rule)
                break
        for pattern in spec.get("regex", []):
            if pattern and re.search(pattern, name):
                matched.append(rule)
                break
    return matched


def detect_header(
    sheet, max_rows: int, header_aliases: Dict[str, List[str]]
) -> Tuple[Optional[int], Dict[str, int]]:
    best_row = None
    best_map = {}
    best_score = 0

    max_scan = min(max_rows, sheet.max_row)
    for offset, row_cells in enumerate(
        sheet.iter_rows(min_row=1, max_row=max_scan, values_only=True),
        start=1,
    ):
        col_map = {}
        score = 0

        for col_idx, cell in enumerate(row_cells, 1):
            value = "" if cell is None else str(cell).strip().lower()
            for key, aliases in header_aliases.items():
                if any(alias in value for alias in aliases):
                    if key not in col_map:
                        col_map[key] = col_idx
                        score += 1

        has_file = "file" in col_map
        has_line = "line" in col_map
        has_source = "source_location" in col_map or "ext_attr1" in col_map
        if score > best_score and ((has_file and has_line) or has_source):
            best_score = score
            best_row = offset
            best_map = col_map

    return best_row, best_map


def collect_uwa_records(
    files: List[Path],
    header_row_override: Optional[int],
    column_override: Dict[str, Optional[int]],
    sheet_name: Optional[str],
    max_header_scan: int,
    rule_excel_map: Dict[str, Dict[str, List[str]]],
    prefer_rule_from_file: bool,
) -> List[ViolationRecord]:
    header_aliases = {
        "file": ["file", "path", "文件", "文件名", "文件路径", "脚本", "资源", "pathname"],
        "line": ["line", "line number", "行号", "行", "行数"],
        "rule": ["rule", "规则", "类型", "违规类型", "问题类型", "分类", "检查项"],
        "desc": ["description", "desc", "说明", "描述", "问题", "详情", "原因", "content"],
        "source_location": [
            "sourcelocation",
            "source location",
            "source_location",
            "source loc",
            "资源位置",
            "源位置",
        ],
        "ext_attr1": ["ext_attr1", "extattr1", "ext attr1"],
    }

    records: List[ViolationRecord] = []
    for file_path in files:
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        sheets = [wb[sheet_name]] if sheet_name and sheet_name in wb.sheetnames else wb.worksheets

        for sheet in sheets:
            header_row = header_row_override
            col_map = {}

            if header_row is None:
                header_row, col_map = detect_header(sheet, max_header_scan, header_aliases)
            else:
                col_map = {}

            for key, col_idx in column_override.items():
                if col_idx:
                    col_map[key] = col_idx

            if not header_row:
                continue

            for row_idx, row in enumerate(
                sheet.iter_rows(
                    min_row=header_row + 1,
                    max_row=sheet.max_row,
                    values_only=True,
                ),
                start=header_row + 1,
            ):

                def cell_value(col_key: str) -> str:
                    col = col_map.get(col_key)
                    if not col:
                        return ""
                    if col - 1 >= len(row):
                        return ""
                    value = row[col - 1]
                    return "" if value is None else str(value).strip()

                source_val = cell_value("source_location")
                if not source_val:
                    source_val = cell_value("ext_attr1")
                rule_val = cell_value("rule")
                desc_val = cell_value("desc")
                if not rule_val and desc_val:
                    rule_val = desc_val

                source_records = parse_source_location(
                    source_val, rule_val, desc_val, str(file_path), sheet.title, row_idx
                )
                if source_records:
                    if rule_excel_map and prefer_rule_from_file:
                        file_rules = rule_from_excel_filename(str(file_path), rule_excel_map)
                        if file_rules:
                            for record in source_records:
                                for file_rule in file_rules:
                                    records.append(
                                        ViolationRecord(
                                            file=record.file,
                                            line=record.line,
                                            rule=file_rule,
                                            description=record.description,
                                            source=record.source,
                                            sheet=record.sheet,
                                            row=record.row,
                                        )
                                    )
                            continue
                    records.extend(source_records)
                    continue

                file_val = cell_value("file")
                if not file_val:
                    continue

                line_val = cell_value("line")
                record = ViolationRecord(
                    file=file_val,
                    line=parse_line_number(line_val),
                    rule=rule_val,
                    description=desc_val,
                    source=str(file_path),
                    sheet=sheet.title,
                    row=row_idx,
                )
                if rule_excel_map and prefer_rule_from_file and not record.rule:
                    file_rules = rule_from_excel_filename(str(file_path), rule_excel_map)
                    if file_rules:
                        for file_rule in file_rules:
                            records.append(
                                ViolationRecord(
                                    file=record.file,
                                    line=record.line,
                                    rule=file_rule,
                                    description=record.description,
                                    source=record.source,
                                    sheet=record.sheet,
                                    row=record.row,
                                )
                            )
                        continue
                records.append(record)

    return records


def collect_local_records(path: Path) -> List[ViolationRecord]:
    with open(path, "r", encoding="utf-8") as f:
        data = json.load(f)

    records: List[ViolationRecord] = []
    for file_result in data.get("files", []):
        file_path = file_result.get("file", "")
        for v in file_result.get("violations", []):
            records.append(
                ViolationRecord(
                    file=file_path,
                    line=v.get("line"),
                    rule=v.get("rule", ""),
                    description=v.get("description", ""),
                    source="local_json",
                    line_content=v.get("line_content", ""),
                )
            )
    return records


def match_file_candidates(
    local_path: str, uwa_paths: Dict[str, List[ViolationRecord]], mode: str
) -> List[str]:
    local_norm = normalize_path(local_path)
    local_base = os.path.basename(local_norm)

    candidates = []
    for uwa_path in uwa_paths.keys():
        uwa_norm = normalize_path(uwa_path)
        uwa_base = os.path.basename(uwa_norm)

        if mode == "exact":
            if local_norm == uwa_norm:
                candidates.append(uwa_path)
        elif mode == "basename":
            if local_base and local_base == uwa_base:
                candidates.append(uwa_path)
        else:
            if local_norm and (uwa_norm.endswith(local_norm) or local_norm.endswith(uwa_norm)):
                candidates.append(uwa_path)
            elif local_base and local_base == uwa_base:
                candidates.append(uwa_path)

    return candidates


def map_uwa_file_to_local_key(
    uwa_path: str, local_keys: List[str], mode: str
) -> Tuple[Optional[str], List[str]]:
    uwa_norm = normalize_path(uwa_path)
    candidates = []
    if mode == "exact":
        if uwa_norm in local_keys:
            candidates.append(uwa_norm)
    elif mode == "basename":
        uwa_base = os.path.basename(uwa_norm)
        for key in local_keys:
            if os.path.basename(key) == uwa_base:
                candidates.append(key)
    else:
        for key in local_keys:
            if uwa_norm.endswith(key) or key.endswith(uwa_norm):
                candidates.append(key)
    if not candidates:
        return None, []
    if len(candidates) == 1:
        return candidates[0], candidates
    candidates.sort(key=len, reverse=True)
    return candidates[0], candidates


def apply_canonical_rule(
    record: ViolationRecord,
    rule_map: Dict[str, str],
    rule_excel_map: Dict[str, Dict[str, List[str]]],
    prefer_rule_from_file: bool,
) -> str:
    canon = normalize_rule(record.rule)
    if canon in rule_map:
        canon = rule_map[canon]
    if (not canon or prefer_rule_from_file) and record.source and rule_excel_map:
        file_rules = rule_from_excel_filename(record.source, rule_excel_map)
        if file_rules:
            canon = file_rules[0]
    return canon or normalize_rule(record.rule)


def match_records_by_line(
    local_records: List[ViolationRecord],
    uwa_records: List[ViolationRecord],
    tolerance: int,
) -> Tuple[List[Tuple[ViolationRecord, ViolationRecord, int]], List[ViolationRecord], List[ViolationRecord]]:
    local_sorted = sorted([r for r in local_records if r.line is not None], key=lambda r: r.line)
    uwa_sorted = sorted([r for r in uwa_records if r.line is not None], key=lambda r: r.line)
    used_uwa = set()
    matches: List[Tuple[ViolationRecord, ViolationRecord, int]] = []
    only_local: List[ViolationRecord] = []

    for local in local_sorted:
        best_index = None
        best_diff = None
        for idx, uwa in enumerate(uwa_sorted):
            if idx in used_uwa:
                continue
            diff = abs(local.line - uwa.line)
            if diff <= tolerance and (best_diff is None or diff < best_diff):
                best_diff = diff
                best_index = idx
        if best_index is not None:
            used_uwa.add(best_index)
            matches.append((local, uwa_sorted[best_index], best_diff or 0))
        else:
            only_local.append(local)

    only_uwa = [uwa for idx, uwa in enumerate(uwa_sorted) if idx not in used_uwa]

    # 保留缺失行号的记录为未匹配
    only_local.extend([r for r in local_records if r.line is None])
    only_uwa.extend([r for r in uwa_records if r.line is None])

    return matches, only_local, only_uwa


def is_desc_match(local_desc: str, uwa_desc: str) -> bool:
    local_norm = normalize_desc(local_desc)
    uwa_norm = normalize_desc(uwa_desc)
    if not local_norm or not uwa_norm:
        return True
    return local_norm in uwa_norm or uwa_norm in local_norm


def build_output_csv(path: Path, rows: List[Dict[str, str]]):
    if not rows:
        return
    with open(path, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def main():
    parser = argparse.ArgumentParser(description="对比UWA Excel结果和本地JSON结果")
    parser.add_argument("--uwa-dir", help="UWA Excel目录")
    parser.add_argument("--uwa-files", nargs="*", help="UWA Excel文件列表")
    parser.add_argument("--local-json", required=True, help="本地检测JSON路径")
    parser.add_argument("-o", "--output", default="uwa_compare_result.json", help="输出JSON文件")
    parser.add_argument("--line-tolerance", type=int, default=3, help="行号容差(默认3)")
    parser.add_argument("--match-by", choices=["rule", "rule+desc"], default="rule")
    parser.add_argument("--rule-map", help="规则映射JSON文件")
    parser.add_argument("--rule-excel-map", help="规则与Excel文件名映射JSON")
    parser.add_argument(
        "--prefer-rule-from-file",
        action="store_true",
        help="优先使用Excel文件名映射规则",
    )
    parser.add_argument("--file-match", choices=["exact", "suffix", "basename"], default="suffix")
    parser.add_argument("--sheet", help="只读取指定sheet")
    parser.add_argument("--max-header-scan", type=int, default=10)
    parser.add_argument("--header-row", type=int, help="固定表头行(1-based)")
    parser.add_argument("--file-col", type=int, help="文件列(1-based)")
    parser.add_argument("--line-col", type=int, help="行号列(1-based)")
    parser.add_argument("--rule-col", type=int, help="规则列(1-based)")
    parser.add_argument("--desc-col", type=int, help="描述列(1-based)")
    parser.add_argument("--source-col", type=int, help="sourceLocation列(1-based)")
    parser.add_argument("--ext-attr1-col", type=int, help="ext_attr1列(1-based)")
    parser.add_argument("--target-path", help="仅对比指定路径下的文件")
    parser.add_argument("--export-csv", action="store_true", help="额外导出CSV")

    args = parser.parse_args()

    uwa_files: List[Path] = []
    if args.uwa_files:
        uwa_files.extend(Path(p) for p in args.uwa_files)
    if args.uwa_dir:
        for root, _, files in os.walk(args.uwa_dir):
            for name in files:
                if name.lower().endswith((".xlsx", ".xlsm")):
                    uwa_files.append(Path(root) / name)

    if not uwa_files:
        print("未找到UWA Excel文件")
        return

    column_override = {
        "file": args.file_col,
        "line": args.line_col,
        "rule": args.rule_col,
        "desc": args.desc_col,
        "source_location": args.source_col,
        "ext_attr1": args.ext_attr1_col,
    }

    rule_map = load_rule_map(args.rule_map)
    rule_excel_map = load_rule_excel_map(args.rule_excel_map)
    if not rule_excel_map:
        rule_excel_map = default_rule_excel_map()
    prefer_rule_from_file = args.prefer_rule_from_file or bool(rule_excel_map)

    uwa_records = collect_uwa_records(
        uwa_files,
        args.header_row,
        column_override,
        args.sheet,
        args.max_header_scan,
        rule_excel_map,
        prefer_rule_from_file,
    )
    local_records = collect_local_records(Path(args.local_json))
    if args.target_path:
        uwa_records = filter_records_by_path(uwa_records, args.target_path)
        local_records = filter_records_by_path(local_records, args.target_path)

    for record in local_records:
        record.canon_rule = apply_canonical_rule(
            record, rule_map, rule_excel_map, prefer_rule_from_file
        )
    for record in uwa_records:
        record.canon_rule = apply_canonical_rule(
            record, rule_map, rule_excel_map, prefer_rule_from_file
        )

    local_keys = sorted({normalize_path(record.file) for record in local_records if record.file})
    display_path_by_key: Dict[str, str] = {}
    for record in local_records:
        key = normalize_path(record.file)
        if key and key not in display_path_by_key:
            display_path_by_key[key] = record.file

    ambiguous = []
    uwa_file_key_cache: Dict[str, Optional[str]] = {}
    for record in uwa_records:
        if record.file in uwa_file_key_cache:
            continue
        mapped_key, candidates = map_uwa_file_to_local_key(
            record.file, local_keys, args.file_match
        )
        uwa_file_key_cache[record.file] = mapped_key
        if len(candidates) > 1:
            ambiguous.append({"uwa_file": record.file, "candidates": candidates})

    rule_results: Dict[str, Dict[str, object]] = {}
    flat_diffs: List[Dict[str, object]] = []
    all_rules = sorted(
        {record.canon_rule for record in local_records + uwa_records if record.canon_rule}
    )

    for rule in all_rules:
        local_rule_records = [r for r in local_records if r.canon_rule == rule]
        uwa_rule_records = [r for r in uwa_records if r.canon_rule == rule]

        local_by_file: Dict[str, List[ViolationRecord]] = {}
        for record in local_rule_records:
            key = normalize_path(record.file)
            local_by_file.setdefault(key, []).append(record)

        uwa_by_file: Dict[str, List[ViolationRecord]] = {}
        for record in uwa_rule_records:
            key = uwa_file_key_cache.get(record.file)
            if not key:
                key = normalize_path(record.file)
                if key not in display_path_by_key:
                    display_path_by_key[key] = record.file
            uwa_by_file.setdefault(key, []).append(record)

        files = sorted(set(list(local_by_file.keys()) + list(uwa_by_file.keys())))
        rule_summary = {"matched": 0, "only_local": 0, "only_uwa": 0}
        file_results: Dict[str, Dict[str, object]] = {}

        for file_key in files:
            local_list = local_by_file.get(file_key, [])
            uwa_list = uwa_by_file.get(file_key, [])
            matches, only_local, only_uwa = match_records_by_line(
                local_list, uwa_list, args.line_tolerance
            )

            rule_summary["matched"] += len(matches)
            rule_summary["only_local"] += len(only_local)
            rule_summary["only_uwa"] += len(only_uwa)

            display_path = display_path_by_key.get(file_key, file_key)
            file_results[display_path] = {
                "matches": [
                    {
                        "local_line": m[0].line,
                        "uwa_line": m[1].line,
                        "line_diff": m[2],
                    }
                    for m in matches
                ],
                "only_local": [
                    {
                        "line": r.line,
                        "line_content": r.line_content,
                        "description": r.description,
                    }
                    for r in only_local
                ],
                "only_uwa": [
                    {
                        "line": r.line,
                        "description": r.description,
                        "source": r.source,
                        "sheet": r.sheet,
                        "row": r.row,
                    }
                    for r in only_uwa
                ],
            }

            for r in only_local:
                flat_diffs.append(
                    {
                        "rule": rule,
                        "file": display_path,
                        "line": r.line,
                        "line_content": r.line_content,
                        "side": "only_json",
                    }
                )
            for r in only_uwa:
                flat_diffs.append(
                    {
                        "rule": rule,
                        "file": display_path,
                        "line": r.line,
                        "side": "only_excel",
                        "source": r.source,
                        "sheet": r.sheet,
                        "row": r.row,
                    }
                )

        rule_results[rule] = {
            "summary": rule_summary,
            "files": file_results,
        }

    output = {
        "summary": {
            "local_total": len(local_records),
            "uwa_total": len(uwa_records),
            "rules": len(rule_results),
            "ambiguous_files": len(ambiguous),
            "target_path": args.target_path or "",
        },
        "rules": rule_results,
        "diffs": flat_diffs,
        "ambiguous": ambiguous,
    }

    with open(args.output, "w", encoding="utf-8") as f:
        json.dump(output, f, ensure_ascii=False, indent=2)

    if args.export_csv:
        base = Path(args.output).with_suffix("")
        build_output_csv(base.with_name(base.name + "_diffs.csv"), flat_diffs)

    print(f"对比完成，结果已保存到: {args.output}")


if __name__ == "__main__":
    main()
