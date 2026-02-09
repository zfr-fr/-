#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Set, Tuple

try:
    import openpyxl
except ImportError:
    sys.stderr.write(
        "Missing dependency: openpyxl. Install with: python3 -m pip install openpyxl\n"
    )
    sys.exit(1)


DEFAULT_ANCHORS = ["Assets/"]
LINE_NUMBER_RE = re.compile(r"\d+")
WINDOWS_ABS_RE = re.compile(r"^[A-Za-z]:[\\/]")


DEFAULT_RULE_KEYWORDS: Dict[str, List[str]] = {
    "DEBUG_LOG_USAGE": [
        "unityengine.debug",
        "debug.log",
        "logerror",
        "logexception",
        "log函数",
        "debug类",
    ],
    "NEW_IN_UPDATE_METHOD": [
        "update",
        "lateupdate",
        "fixedupdate",
        "引用类型",
        "new操作",
    ],
    "EMPTY_UPDATE_METHOD": [
        "空的update",
        "空update",
        "lateupdate",
        "fixedupdate",
    ],
    "FIND_FUNCTION_CALLS": [
        "findobjectoftype",
        "findobjectsoftype",
        "findgameobjectwithtag",
        "findgameobjectswithtag",
        "findobjectsoftypeall",
        "find类函数",
    ],
    "FOREACH_ON_SPECIFIC_CONTAINERS": [
        "foreach",
        "foreach遍历",
        "sorteddictionary",
        "hashtable",
        "bitarray",
        "queue",
        "sortedlist",
        "arraylist",
        "stack",
        "特定容器",
    ],
    "GET_COMPONENTS_IN_CHILDREN": ["getcomponentsinchildren"],
    "GET_COMPONENTS_IN_PARENT": ["getcomponentsinparent"],
    "COMPUTE_BUFFER_GETDATA": ["computebuffer.getdata", "computebuffer", "getdata"],
    "TEXTURE_GETPIXELS_CALL": ["getpixels32", "getpixels", "texture.getpixels"],
    "TEXTASSET_BYTES_USAGE": ["textasset.bytes", "textasset", "bytes"],
    "CAMERA_MAIN_USAGE": ["camera.main"],
    "RENDERER_MATERIAL_GETTER": [
        "material_materials",
        "renderer",
        "material",
        "materials",
    ],
    "RENDERER_SHAREDMATERIALS_GETTER": ["sharedmaterials", "renderer"],
    "ONGUI_METHOD_PRESENT": ["ongui"],
    "SENDMESSAGE_USAGE": ["sendmessage"],
    "TEXTURE_SETPIXELS_CALL": ["setpixels", "texture.setpixels"],
    "TAG_PROPERTY_USAGE": ["tag属性", ".tag", "tag"],
    "LINQ_USAGE": ["linq"],
    "REFLECTION_USAGE": ["reflection", "反射", "getinterface"],
}


@dataclass
class ScopeStats:
    total: int = 0
    in_scope: int = 0
    out_scope: int = 0
    out_samples: List[str] = field(default_factory=list)

    def record(self, entry: str, in_scope: bool, limit: int) -> None:
        self.total += 1
        if in_scope:
            self.in_scope += 1
            return
        self.out_scope += 1
        if limit > 0 and len(self.out_samples) < limit:
            self.out_samples.append(entry)


@dataclass
class ExcelSheetInfo:
    sheet: str
    header_row: Optional[int]
    column_index: Optional[int]
    parsed_entries: int


def normalize_slashes(value: str) -> str:
    return value.replace("\\", "/")


def is_windows_abs(path: str) -> bool:
    return bool(WINDOWS_ABS_RE.match(path))


def is_abs(path: str) -> bool:
    return path.startswith("/") or is_windows_abs(path)


class PathNormalizer:
    def __init__(
        self,
        project_root: Path,
        strip_prefixes: Sequence[str],
        anchors: Sequence[str],
    ) -> None:
        self.project_root = project_root
        self.project_root_norm = normalize_slashes(str(project_root)).rstrip("/")
        self.project_root_norm_lower = self.project_root_norm.lower()
        self.strip_prefixes = [
            normalize_slashes(prefix).rstrip("/")
            for prefix in strip_prefixes
            if prefix
        ]
        self.anchors = [normalize_slashes(anchor).strip("/") for anchor in anchors if anchor]
        self.anchor_tails = self._build_anchor_tails()
        self.anchor_tail = self._select_anchor_tail()
        self.display_by_key: Dict[str, str] = {}
        self.candidates_by_key: Dict[str, List[Path]] = {}

    def _build_anchor_tails(self) -> List[str]:
        tails: List[str] = []
        for anchor in self.anchors:
            anchor_lower = anchor.lower()
            index = self.project_root_norm_lower.find(anchor_lower)
            if index != -1:
                tail = self.project_root_norm[index:].rstrip("/")
                tails.append(tail)
        return tails

    def _select_anchor_tail(self) -> Optional[str]:
        if not self.anchor_tails:
            return None
        return max(self.anchor_tails, key=len)

    def is_in_scope(self, raw_path: str) -> bool:
        path = normalize_slashes(raw_path.strip().strip('"')).rstrip("/")
        if not path:
            return False
        path_lower = path.lower()
        if is_abs(path):
            if path_lower == self.project_root_norm_lower:
                return True
            return path_lower.startswith(self.project_root_norm_lower + "/")
        for anchor in self.anchors:
            anchor_lower = anchor.lower()
            if path_lower == anchor_lower or path_lower.startswith(anchor_lower + "/"):
                if not self.anchor_tails:
                    return False
                for tail in self.anchor_tails:
                    tail_lower = tail.lower()
                    if path_lower == tail_lower or path_lower.startswith(tail_lower + "/"):
                        return True
                return False
        return True

    def normalize_display(self, raw_path: str) -> str:
        path = normalize_slashes(raw_path.strip().strip('"'))
        if is_abs(path):
            path_norm = path.rstrip("/")
            path_lower = path_norm.lower()
            if path_lower == self.project_root_norm_lower:
                if self.anchor_tail:
                    return self.anchor_tail
                return ""
            if path_lower.startswith(self.project_root_norm_lower + "/"):
                rel = path_norm[len(self.project_root_norm) + 1 :]
                if self.anchor_tail:
                    return f"{self.anchor_tail}/{rel}"
                return rel
        for prefix in self.strip_prefixes:
            prefix_lower = prefix.lower()
            if path.lower().startswith(prefix_lower + "/"):
                path = path[len(prefix) + 1 :]
                break
        for anchor in self.anchors:
            anchor_lower = anchor.lower()
            index = path.lower().find(anchor_lower)
            if index != -1:
                path = path[index:]
                break
        return path

    def register(self, raw_path: str) -> str:
        display = self.normalize_display(raw_path)
        key = display.lower()
        if key not in self.display_by_key:
            self.display_by_key[key] = display
        self._add_candidates(key, raw_path, display)
        return key

    def _add_candidates(self, key: str, raw_path: str, display: str) -> None:
        candidates = self.candidates_by_key.setdefault(key, [])

        def add(path_str: str) -> None:
            if not path_str:
                return
            path = Path(path_str)
            if path not in candidates:
                candidates.append(path)

        raw_norm = normalize_slashes(raw_path)
        if is_abs(raw_norm):
            add(raw_norm)
        if display:
            add(str(self.project_root / display))
            if is_abs(display):
                add(display)
        if raw_norm and not is_abs(raw_norm):
            add(str(self.project_root / raw_norm))

    def display_for_key(self, key: str) -> str:
        return self.display_by_key.get(key, key)

    def resolve_path(self, key: str) -> Optional[Path]:
        for candidate in self.candidates_by_key.get(key, []):
            try:
                if candidate.exists():
                    return candidate
            except OSError:
                continue
        return None


class FileContentCache:
    def __init__(self, resolver: PathNormalizer) -> None:
        self.resolver = resolver
        self._cache: Dict[Path, List[str]] = {}

    def get_line(self, key: str, line_number: int) -> Tuple[str, Optional[str]]:
        path = self.resolver.resolve_path(key)
        if not path:
            return "<file not found>", None
        if path not in self._cache:
            try:
                content = path.read_text(encoding="utf-8", errors="replace")
            except OSError:
                return "<file unreadable>", str(path)
            self._cache[path] = content.splitlines()
        lines = self._cache[path]
        if line_number < 1 or line_number > len(lines):
            return "<line out of range>", str(path)
        return lines[line_number - 1].rstrip("\n"), str(path)


def parse_line_numbers(value: object) -> List[int]:
    if value is None:
        return []
    if isinstance(value, int):
        return [value]
    if isinstance(value, list):
        numbers = []
        for item in value:
            numbers.extend(parse_line_numbers(item))
        return numbers
    numbers = [int(n) for n in LINE_NUMBER_RE.findall(str(value))]
    return numbers


def parse_ext_attr1(value: object) -> Iterable[Tuple[str, int]]:
    if value is None:
        return []
    text = str(value).replace("\n", ";").replace("\r", ";")
    results: List[Tuple[str, int]] = []
    for segment in text.split(";"):
        segment = segment.strip()
        if not segment or ":" not in segment:
            continue
        path_part, line_part = segment.rsplit(":", 1)
        path_part = path_part.strip()
        if not path_part:
            continue
        for number in parse_line_numbers(line_part):
            results.append((path_part, number))
    return results


def iter_json_violations(data: object) -> Iterable[Tuple[str, str, int]]:
    entries: List[object]
    if isinstance(data, dict) and isinstance(data.get("files"), list):
        entries = data["files"]
    elif isinstance(data, list):
        entries = data
    else:
        entries = [data]

    for entry in entries:
        if not isinstance(entry, dict):
            continue
        entry_file = entry.get("file") or entry.get("path")
        violations = entry.get("violations")
        if isinstance(violations, list):
            for violation in violations:
                if not isinstance(violation, dict):
                    continue
                file_path = violation.get("file") or violation.get("path") or entry_file
                rule = violation.get("rule") or violation.get("rule_code")
                if isinstance(rule, str):
                    rule = rule.strip()
                line_value = (
                    violation.get("line")
                    or violation.get("line_number")
                    or violation.get("line_no")
                )
                for number in parse_line_numbers(line_value):
                    if file_path and rule:
                        yield str(file_path), str(rule), number
        else:
            rule = entry.get("rule") or entry.get("rule_code")
            if isinstance(rule, str):
                rule = rule.strip()
            line_value = entry.get("line") or entry.get("line_number") or entry.get("line_no")
            for number in parse_line_numbers(line_value):
                if entry_file and rule:
                    yield str(entry_file), str(rule), number


def find_ext_attr1_column(sheet, max_scan_rows: Optional[int]) -> Optional[Tuple[int, int]]:
    max_rows = sheet.max_row or 0
    if max_scan_rows and max_scan_rows > 0:
        max_rows = min(max_rows, max_scan_rows)
    for row_index in range(1, max_rows + 1):
        row = sheet[row_index]
        for cell_index, cell in enumerate(row, start=1):
            value = "" if cell.value is None else str(cell.value).strip()
            if "ext_attr1" in value.lower():
                return cell_index, row_index
    return None


def load_excel_violations(
    excel_path: Path,
    normalizer: PathNormalizer,
    scope_stats: Optional[ScopeStats] = None,
    scope_limit: int = 0,
    header_scan_rows: Optional[int] = None,
    sheet_infos: Optional[List[ExcelSheetInfo]] = None,
) -> Set[Tuple[str, int]]:
    workbook = openpyxl.load_workbook(excel_path, data_only=True, read_only=True)
    try:
        results: Set[Tuple[str, int]] = set()
        any_column_found = False
        for sheet in workbook.worksheets:
            column_info = find_ext_attr1_column(sheet, header_scan_rows)
            if not column_info:
                if sheet_infos is not None:
                    sheet_infos.append(
                        ExcelSheetInfo(sheet=sheet.title, header_row=None, column_index=None, parsed_entries=0)
                    )
                continue
            any_column_found = True
            column_index, header_row = column_info
            before_count = len(results)
            for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
                if not row or len(row) < column_index:
                    continue
                value = row[column_index - 1]
                for path_value, line_number in parse_ext_attr1(value):
                    in_scope = normalizer.is_in_scope(path_value)
                    if scope_stats:
                        scope_stats.record(f"{path_value}:{line_number}", in_scope, scope_limit)
                    if not in_scope:
                        continue
                    key = normalizer.register(path_value)
                    results.add((key, line_number))
            parsed_entries = len(results) - before_count
            if sheet_infos is not None:
                sheet_infos.append(
                    ExcelSheetInfo(
                        sheet=sheet.title,
                        header_row=header_row,
                        column_index=column_index,
                        parsed_entries=parsed_entries,
                    )
                )
        if not any_column_found:
            sys.stderr.write(f"No ext_attr1 column found in {excel_path}\n")
        return results
    finally:
        workbook.close()


def load_json_violations(
    json_path: Path,
    normalizer: PathNormalizer,
    scope_stats: Optional[ScopeStats] = None,
    scope_limit: int = 0,
) -> Dict[str, Set[Tuple[str, int]]]:
    with json_path.open("r", encoding="utf-8") as handle:
        data = json.load(handle)
    results: Dict[str, Set[Tuple[str, int]]] = {}
    for file_path, rule, line_number in iter_json_violations(data):
        in_scope = normalizer.is_in_scope(file_path)
        if scope_stats:
            scope_stats.record(f"{file_path}:{line_number}", in_scope, scope_limit)
        if not in_scope:
            continue
        key = normalizer.register(file_path)
        results.setdefault(rule, set()).add((key, line_number))
    return results


def load_rule_map(rule_map_path: Optional[Path]) -> Dict[str, str]:
    if not rule_map_path:
        return {}
    with rule_map_path.open("r", encoding="utf-8") as handle:
        data = json.load(handle)
    excel_to_rule: Dict[str, str] = {}
    if not isinstance(data, dict):
        return excel_to_rule
    if all(isinstance(value, str) for value in data.values()):
        for excel_name, rule in data.items():
            excel_to_rule[str(excel_name).lower()] = str(rule)
    else:
        for rule, excel_names in data.items():
            if isinstance(excel_names, str):
                excel_names = [excel_names]
            if not isinstance(excel_names, list):
                continue
            for name in excel_names:
                excel_to_rule[str(name).lower()] = str(rule)
    return excel_to_rule


def match_rule_for_excel(
    excel_path: Path, excel_to_rule: Dict[str, str]
) -> Optional[str]:
    name_lower = excel_path.name.lower()
    stem_lower = excel_path.stem.lower()
    if name_lower in excel_to_rule:
        return excel_to_rule[name_lower]
    if stem_lower in excel_to_rule:
        return excel_to_rule[stem_lower]

    best_rule: Optional[str] = None
    best_score = 0
    for rule, keywords in DEFAULT_RULE_KEYWORDS.items():
        score = 0
        for keyword in keywords:
            if keyword and keyword.lower() in name_lower:
                score += len(keyword)
        if score > best_score:
            best_score = score
            best_rule = rule
    return best_rule


def build_rule_report(
    rule: str,
    json_set: Set[Tuple[str, int]],
    excel_set: Set[Tuple[str, int]],
    resolver: PathNormalizer,
    content_cache: FileContentCache,
    excel_file: Optional[str],
    differences: List[Dict[str, object]],
) -> Dict[str, object]:
    only_json = sorted(json_set - excel_set, key=lambda item: (resolver.display_for_key(item[0]), item[1]))
    only_excel = sorted(excel_set - json_set, key=lambda item: (resolver.display_for_key(item[0]), item[1]))
    matched = json_set & excel_set

    def build_entries(items: Iterable[Tuple[str, int]], side: str) -> List[Dict[str, object]]:
        entries: List[Dict[str, object]] = []
        for key, line in items:
            display_path = resolver.display_for_key(key)
            line_content, resolved_path = content_cache.get_line(key, line)
            entry = {
                "file": display_path,
                "line": line,
                "content": line_content,
            }
            if resolved_path:
                entry["resolved_path"] = resolved_path
            entries.append(entry)
            differences.append(
                {
                    "rule": rule,
                    "side": side,
                    "file": display_path,
                    "line": line,
                    "content": line_content,
                    "excel_file": excel_file,
                }
            )
        return entries

    presence = "both"
    if excel_set and not json_set:
        presence = "excel_only"
    elif json_set and not excel_set:
        presence = "json_only"

    return {
        "rule": rule,
        "excel_file": excel_file,
        "presence": presence,
        "json_count": len(json_set),
        "excel_count": len(excel_set),
        "matched_count": len(matched),
        "only_in_json_count": len(only_json),
        "only_in_excel_count": len(only_excel),
        "only_in_json": build_entries(only_json, "only_in_json"),
        "only_in_excel": build_entries(only_excel, "only_in_excel"),
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Compare JSON violations against Excel detection data."
    )
    parser.add_argument("--json", required=True, help="Path to JSON report file.")
    parser.add_argument("--excel-dir", required=True, help="Directory containing Excel files.")
    parser.add_argument(
        "--project-root",
        default=".",
        help="Local project root used to read source files.",
    )
    parser.add_argument(
        "--strip-prefix",
        action="append",
        default=[],
        help="Prefix to strip from file paths before comparison. Repeatable.",
    )
    parser.add_argument(
        "--path-anchor",
        action="append",
        default=[],
        help="Anchor substring to trim path before comparison. Repeatable.",
    )
    parser.add_argument(
        "--rule-map",
        help="Optional JSON file mapping Excel names to rule codes.",
    )
    parser.add_argument(
        "--output",
        default="comparison_report.json",
        help="Output JSON report path.",
    )
    parser.add_argument(
        "--excel-glob",
        default="**/*.xlsx",
        help="Glob pattern for Excel files (relative to excel-dir).",
    )
    parser.add_argument(
        "--debug-rule",
        action="append",
        default=[],
        help="Print debug output for specific rule(s). Repeatable.",
    )
    parser.add_argument(
        "--debug-limit",
        type=int,
        default=50,
        help="Max debug entries per section.",
    )
    parser.add_argument(
        "--debug-file",
        help="Only show debug entries whose file contains this substring.",
    )
    parser.add_argument(
        "--list-rules",
        action="store_true",
        help="Print rule counts for JSON and Excel data.",
    )
    parser.add_argument(
        "--debug-scope",
        action="store_true",
        help="Print how many entries were filtered by project-root scope.",
    )
    parser.add_argument(
        "--debug-scope-limit",
        type=int,
        default=20,
        help="Max filtered entry samples to print.",
    )
    parser.add_argument(
        "--debug-excel",
        action="store_true",
        help="Print per-excel sheet parsing details.",
    )
    parser.add_argument(
        "--excel-header-scan",
        type=int,
        default=500,
        help="Max rows to scan for ext_attr1 header (0 for all).",
    )
    return parser.parse_args()


def _format_entry(resolver: PathNormalizer, item: Tuple[str, int]) -> str:
    path_key, line = item
    return f"{resolver.display_for_key(path_key)}:{line}"


def _filter_by_file(
    resolver: PathNormalizer, items: Iterable[Tuple[str, int]], file_filter: Optional[str]
) -> List[Tuple[str, int]]:
    if not file_filter:
        return list(items)
    token = file_filter.lower()
    return [item for item in items if token in resolver.display_for_key(item[0]).lower()]


def main() -> int:
    args = parse_args()
    project_root = Path(args.project_root).expanduser().resolve()
    json_path = Path(args.json).expanduser().resolve()
    excel_dir = Path(args.excel_dir).expanduser().resolve()
    output_path = Path(args.output).expanduser().resolve()
    rule_map_path = Path(args.rule_map).expanduser().resolve() if args.rule_map else None

    if not json_path.exists():
        sys.stderr.write(f"JSON file not found: {json_path}\n")
        return 2
    if not excel_dir.exists():
        sys.stderr.write(f"Excel directory not found: {excel_dir}\n")
        return 2

    anchors = DEFAULT_ANCHORS + args.path_anchor
    strip_prefixes = [str(project_root)] + args.strip_prefix
    normalizer = PathNormalizer(project_root, strip_prefixes, anchors)
    content_cache = FileContentCache(normalizer)

    excel_to_rule = load_rule_map(rule_map_path)
    json_scope_stats = ScopeStats() if args.debug_scope else None
    excel_scope_stats = ScopeStats() if args.debug_scope else None
    scope_limit = max(0, args.debug_scope_limit)
    excel_files = [
        path
        for path in excel_dir.glob(args.excel_glob)
        if path.is_file() and not path.name.startswith("~$")
    ]

    excel_violations: Dict[str, Set[Tuple[str, int]]] = {}
    excel_file_map: Dict[str, str] = {}
    unmapped_excel_files: List[str] = []

    for excel_path in sorted(excel_files):
        rule = match_rule_for_excel(excel_path, excel_to_rule)
        if not rule:
            unmapped_excel_files.append(str(excel_path))
            continue
        excel_file_map[rule] = excel_path.name
        sheet_infos: List[ExcelSheetInfo] = []
        violations = load_excel_violations(
            excel_path,
            normalizer,
            scope_stats=excel_scope_stats,
            scope_limit=scope_limit,
            header_scan_rows=args.excel_header_scan,
            sheet_infos=sheet_infos if args.debug_excel else None,
        )
        excel_violations.setdefault(rule, set()).update(violations)
        if args.debug_excel:
            print(f"\n[DEBUG] Excel file: {excel_path.name} -> {rule}")
            if not sheet_infos:
                print("  No sheets parsed.")
            for info in sheet_infos:
                if info.column_index is None:
                    print(f"  - Sheet {info.sheet}: ext_attr1 not found")
                else:
                    print(
                        f"  - Sheet {info.sheet}: ext_attr1 at row {info.header_row}, "
                        f"col {info.column_index}, parsed {info.parsed_entries}"
                    )

    json_violations = load_json_violations(
        json_path,
        normalizer,
        scope_stats=json_scope_stats,
        scope_limit=scope_limit,
    )

    if args.debug_scope:
        print("\n[DEBUG] Scope filter stats:")
        if json_scope_stats:
            print(
                f"  JSON total={json_scope_stats.total} "
                f"in_scope={json_scope_stats.in_scope} "
                f"out_scope={json_scope_stats.out_scope}"
            )
            for item in json_scope_stats.out_samples:
                print(f"    - JSON filtered: {item}")
        if excel_scope_stats:
            print(
                f"  Excel total={excel_scope_stats.total} "
                f"in_scope={excel_scope_stats.in_scope} "
                f"out_scope={excel_scope_stats.out_scope}"
            )
            for item in excel_scope_stats.out_samples:
                print(f"    - Excel filtered: {item}")

    all_rules = sorted(set(json_violations.keys()) | set(excel_violations.keys()))
    differences: List[Dict[str, object]] = []
    rule_reports: List[Dict[str, object]] = []

    for rule in all_rules:
        rule_report = build_rule_report(
            rule,
            json_violations.get(rule, set()),
            excel_violations.get(rule, set()),
            normalizer,
            content_cache,
            excel_file_map.get(rule),
            differences,
        )
        rule_reports.append(rule_report)

    summary = {
        "rules_total": len(all_rules),
        "rules_with_differences": sum(
            1
            for report in rule_reports
            if report["only_in_json_count"] or report["only_in_excel_count"]
        ),
        "only_in_json_total": sum(report["only_in_json_count"] for report in rule_reports),
        "only_in_excel_total": sum(report["only_in_excel_count"] for report in rule_reports),
        "unmapped_excel_files": unmapped_excel_files,
        "json_rules_without_excel": sorted(
            rule for rule in json_violations.keys() if rule not in excel_violations
        ),
    }

    if args.list_rules:
        print("\n[DEBUG] JSON rules:")
        for rule, items in sorted(json_violations.items()):
            print(f"  - {rule}: {len(items)}")
        print("[DEBUG] Excel rules:")
        for rule, items in sorted(excel_violations.items()):
            print(f"  - {rule}: {len(items)}")

    debug_rules = [rule.strip() for rule in args.debug_rule if rule.strip()]
    if debug_rules:
        limit = max(0, args.debug_limit)
        file_filter = args.debug_file
        if "*" in debug_rules:
            debug_rules = all_rules
        for rule in debug_rules:
            json_set = json_violations.get(rule, set())
            excel_set = excel_violations.get(rule, set())
            json_set = set(_filter_by_file(normalizer, json_set, file_filter))
            excel_set = set(_filter_by_file(normalizer, excel_set, file_filter))
            matched = json_set & excel_set
            only_json = sorted(json_set - excel_set, key=lambda item: (_format_entry(normalizer, item)))
            only_excel = sorted(excel_set - json_set, key=lambda item: (_format_entry(normalizer, item)))
            matched_sorted = sorted(matched, key=lambda item: (_format_entry(normalizer, item)))

            print(f"\n[DEBUG] Rule: {rule}")
            print(f"  JSON count: {len(json_set)}")
            print(f"  Excel count: {len(excel_set)}")
            print(f"  Matched count: {len(matched)}")
            if file_filter:
                print(f"  File filter: {file_filter}")
            if only_json:
                print("  JSON only:")
                for item in only_json[:limit]:
                    print(f"    - {_format_entry(normalizer, item)}")
            if only_excel:
                print("  Excel only:")
                for item in only_excel[:limit]:
                    print(f"    - {_format_entry(normalizer, item)}")
            if matched_sorted and limit > 0:
                print("  Matched:")
                for item in matched_sorted[:limit]:
                    print(f"    - {_format_entry(normalizer, item)}")

    report = {
        "generated_at": datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S"),
        "project_root": str(project_root),
        "json_file": str(json_path),
        "excel_dir": str(excel_dir),
        "rules": rule_reports,
        "differences": differences,
        "summary": summary,
    }

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8") as handle:
        json.dump(report, handle, ensure_ascii=False, indent=2)

    print(f"Report written to: {output_path}")
    print(f"Rules total: {summary['rules_total']}")
    print(f"Differences (JSON only): {summary['only_in_json_total']}")
    print(f"Differences (Excel only): {summary['only_in_excel_total']}")
    if summary["unmapped_excel_files"]:
        print(f"Unmapped Excel files: {len(summary['unmapped_excel_files'])}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
